#!/usr/bin/env python3
"""ExcelのCF表シートから計算済みの値を取り出し、test/fixtures/excel-cf.json に書く。

使い方:
    python3 -I tools/extract_excel_fixture.py <ライフプラン表.xlsx>

Excelをサンプル入力のまま保存したファイルを渡すこと（アプリの初期データ app/sample.js と同じ入力）。
Excel本体はリポジトリに入れない（公開リポジトリのため）。取り出すのは数値だけ。
"""
import json
import sys
from pathlib import Path

import openpyxl

# CF表シートの行番号 -> テストで使うキー
ROWS = {
    10: 'incomeTotal',
    11: 'salary1',
    12: 'salary2',
    14: 'pension',
    15: 'nisaOut',
    16: 'idecoOut',
    18: 'expenseTotal',
    19: 'living',
    20: 'housing',
    21: 'education',
    22: 'insurance',
    23: 'car',
    24: 'other',
    25: 'nisaIn',
    26: 'idecoIn',
    28: 'balance',
    30: 'total',
    31: 'deposit',
    32: 'nisa',
    33: 'ideco',
}
FIRST_COL = 4  # D列が開始年


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['CF表']
    years = []
    col = FIRST_COL
    while True:
        year = ws.cell(2, col).value
        age = ws.cell(3, col).value
        if not isinstance(year, (int, float)) or age in (None, ''):
            break
        row = {'year': int(year)}
        for r, key in ROWS.items():
            v = ws.cell(r, col).value
            row[key] = float(v) if isinstance(v, (int, float)) else 0.0
        years.append(row)
        col += 1
    if not years:
        sys.exit('CF表から年が読み取れませんでした（計算済みの値が保存されていない可能性があります。Excelで開いて保存し直してください）')
    out = Path(__file__).resolve().parent.parent / 'test' / 'fixtures' / 'excel-cf.json'
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps({'source': Path(path).name, 'years': years}, ensure_ascii=False, indent=1))
    print(f'{len(years)}年分（{years[0]["year"]}〜{years[-1]["year"]}）を {out} に書き出しました')


if __name__ == '__main__':
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
