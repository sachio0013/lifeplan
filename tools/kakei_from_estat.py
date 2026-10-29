"""総務省「家計調査」（e-Stat）の表から、app/book.js の KAKEI（基本生活費などの平均）を作る。

使い方:
  python3 tools/kakei_from_estat.py 年間収入階級別.xlsx 世帯人員世帯主年齢別_1月.xlsx ... 世帯人員世帯主年齢別_12月.xlsx

- 年間収入階級別：家計収支編 二人以上の世帯 詳細結果表 年次「第2-3表 年間収入階級別」（シート「二人・全国」）
- 世帯人員・世帯主の年齢階級別：家計収支編 二人以上の世帯 詳細結果表 月次「第3-13表」を12か月分（年次の表には無いため平均する）

作り方（excelcf.net「年齢層別2人・3人・4人・5人以上家族の生活費と世帯年収別生活費」と同じ）:
  1か月の支出×12。基本生活費＝消費支出−住居−教育−自動車等関係費、住居費＝住居、教育費＝教育、自動車＝自動車等関係費（万円に四捨五入）。
  保険料は家計調査に無いので、生命保険文化センター「生命保険に関する全国実態調査」の値を book.js に手で入れる。
"""
import sys
import openpyxl

KEYS = ['消費支出', '住居', '教育', '自動車等関係費']


def table(path, sheet=None):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb[sheet or wb.sheetnames[0]].iter_rows(values_only=True))
    hdr = rows[8]
    data = {}
    for r in rows[9:]:
        if r[11] and r[12] == '円':
            data.setdefault(r[11], r)   # 同じ名前が下にも出る（交際費の内訳など）ので、最初の行
    cols = [(j, h) for j, h in enumerate(hdr) if j >= 13 and h]
    return {h: {k: float(data[k][j]) for k in KEYS} for j, h in cols}


def yearly(v):
    man = lambda x: round(x * 12 / 10000)
    return [man(v['消費支出'] - v['住居'] - v['教育'] - v['自動車等関係費']), man(v['住居']), man(v['教育']), man(v['自動車等関係費'])]


def main(income, months):
    print('// 年間収入階級別 [基本生活費, 住居費, 教育費, 自動車]')
    for h, v in table(income, '二人・全国').items():
        print(h, yearly(v))
    ts = [table(m) for m in months]
    print('// 世帯人員・世帯主の年齢階級別（%dか月の平均） [基本生活費, 住居費, 教育費, 自動車]' % len(ts))
    for h in ts[0]:
        avg = {k: sum(t[h][k] for t in ts) / len(ts) for k in KEYS}
        print(h, yearly(avg))


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2:])
