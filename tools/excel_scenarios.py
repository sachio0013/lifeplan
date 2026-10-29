#!/usr/bin/env python3
"""アプリの入力（test/scenarios/*.json）をExcelに流し込み、LibreOfficeで再計算して、
CF表の値を test/fixtures/scenario-<名前>.json に書く。npm test がアプリの計算と照合する。

使い方:
    python3 -I tools/excel_scenarios.py <ライフプラン表.xlsx> [作業フォルダ]

必要なもの: openpyxl、LibreOffice（libreoffice-calc）。

LibreOfficeはExcelの給与の数式にある MAX(INDEX(配列,0)) を計算できないため、作業用のコピーでだけ
同じ意味の SUMPRODUCT(MAX(配列)) に置き換える。置き換えたうえでサンプル入力を再計算し、Excelが保存した
値と一致することを最初に確かめる（一致しなければ止まる）。
"""
import datetime
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
CF_ROWS = {
    10: 'incomeTotal', 11: 'salary1', 12: 'salary2', 13: 'salary3', 14: 'pension', 15: 'nisaOut', 16: 'idecoOut',
    18: 'expenseTotal', 19: 'living', 20: 'housing', 21: 'education', 22: 'insurance', 23: 'car', 24: 'other',
    25: 'nisaIn', 26: 'idecoIn', 28: 'balance', 30: 'total', 31: 'deposit', 32: 'nisa', 33: 'ideco',
}
SALARY_BLOCKS = [(33, 8), (44, 8), (55, 5)]  # (先頭行, 行数)
EDU_BLOCKS = [137, 146, 155, 164]


def patch_for_libreoffice(wb):
    pat = re.compile(r'MAX\(INDEX\(\(([^()]*?)\)\*([^(),]*?),0\)\)')
    ws = wb['CF表']
    n = 0
    for row in ws.iter_rows(min_row=11, max_row=13):
        for c in row:
            if isinstance(c.value, str) and 'MAX(INDEX(' in c.value:
                c.value = pat.sub(r'SUMPRODUCT(MAX((\1)*\2))', c.value)
                n += 1
    return n


def birth_year(state, member_id):
    for m in state['family']:
        if m['id'] == member_id:
            return int(m['birth'][:4])
    return None


def put_rows(ws, first, count, cols, values):
    """first行から count 行を空にしてから values（タプルの列）を書く。cols は列名のタプル。"""
    for i in range(count):
        for c in cols:
            ws[f'{c}{first + i}'] = None
    for i, v in enumerate(values[:count]):
        for c, x in zip(cols, v):
            ws[f'{c}{first + i}'] = x


def write_state(wb, s):
    ws = wb['入力シート']
    sy = s['startYear']
    ws['C12'] = datetime.datetime(sy, 1, 1)
    ws['N12'] = s['endAge']
    ws['H12'], ws['H13'], ws['H14'] = (a['amount'] for a in s['assets'][:3])
    wb['CF表']['C31'] = s.get('depositRate', 0) / 100

    for i in range(6):
        r = 19 + i
        m = s['family'][i] if i < len(s['family']) else None
        ws[f'D{r}'] = m['name'] if m else None
        ws[f'I{r}'] = m['relation'] if m else None
        ws[f'L{r}'] = datetime.datetime(*map(int, m['birth'].split('/'))) if m else None

    # 給与（Excelは退職を「0の行」で表すので、退職年齢から0の行を足す）
    for b, (first, count) in enumerate(SALARY_BLOCKS):
        sal = s['salaries'][b] if b < len(s['salaries']) else None
        vals = []
        if sal:
            rows = sorted(sal['rows'], key=lambda r: r['year'])
            for k, r in enumerate(rows):
                vals.append((sy if k == 0 else r['year'], r['amount'], r['growth'] / 100))
            by = birth_year(s, sal.get('memberId'))
            if sal.get('retireAge') and by:
                vals.append((by + sal['retireAge'], 0, 0))
        put_rows(ws, first, count, ('B', 'F', 'L'), vals)
        for i in range(count):
            if ws[f'L{first + i}'].value is None:
                ws[f'L{first + i}'] = 0

    put_rows(ws, 64, 3, ('B', 'F'), [(r['year'], r['amount']) for r in s['lump']])
    put_rows(ws, 69, 5, ('B', 'F', 'J'), [(r['year'], r['toYear'], r['amount']) for r in s['corpPension']])
    put_rows(ws, 77, 5, ('B', 'F', 'J'), [(r['year'], r['toYear'], r['amount']) for r in s['privPension']])
    put_rows(ws, 86, 5, ('B', 'F'), [(r['year'], r['amount']) for r in s['pubPension']])

    lv = s['living']
    ws['B107'], ws['F107'], ws['L107'] = sy, lv[0]['amount'], s['inflation'] / 100
    put_rows(ws, 109, 4, ('B', 'F'), [(r['year'], r['amount']) for r in lv[1:]])
    hs = s['housing']
    ws['B115'], ws['F115'] = sy, hs[0]['amount']
    put_rows(ws, 117, 4, ('B', 'F'), [(r['year'], r['amount']) for r in hs[1:]])
    put_rows(ws, 123, 3, ('B', 'F'), [(r['year'], r['amount']) for r in s['housingLump']])

    kids = [m for m in s['family'] if m['relation'].startswith('子供')]
    for b, base in enumerate(EDU_BLOCKS):
        kid = kids[b] if b < len(kids) else None
        ws[f'F{base}'] = datetime.datetime(*map(int, kid['birth'].split('/'))) if kid else None
        if not kid:
            continue
        picks = s['edu'].get(str(kid['id']), [])
        for i, stage in enumerate(picks):
            ws[f'B{base + 1 + i}'] = stage
        for i, span in s['eduSpan'].get(str(kid['id']), {}).items():
            ws[f'D{base + 1 + int(i)}'], ws[f'H{base + 1 + int(i)}'] = span['from'], span['to']
        for i, amt in s['eduAmount'].get(str(kid['id']), {}).items():
            ws[f'L{base + 1 + int(i)}'] = amt

    put_rows(ws, 184, 7, ('B', 'F'), [(r['year'], r['amount']) for r in s['insurance']])
    put_rows(ws, 195, 5, ('B', 'F'), [(r['year'], r['amount']) for r in s['carCost']])
    c = s['carBuy']
    ws['B202'], ws['F202'], ws['J202'], ws['N202'] = c['fromYear'], c['every'], c['amount'], c['untilAge']
    for k in range(16):
        y = c['fromYear'] + k * c['every']
        if str(y) in s['carBuyOverride']:
            ws[f'E{204 + k}'] = s['carBuyOverride'][str(y)]

    put_rows(ws, 231, 5, ('B', 'F'), [(r['year'], r['amount']) for r in s['otherFlow']])
    put_rows(ws, 238, 5, ('B', 'F'), [(r['year'], r['amount']) for r in s['otherLump']])

    for asset, rate_row, add_row, out_row in ((s['assets'][1], 257, 258, 265), (s['assets'][2], 272, 273, 280)):
        ws[f'J{rate_row}'] = asset['rate'] / 100
        ws[f'L{rate_row}'] = float(asset['tax'].rstrip('%')) / 100
        iv = s['invest'][str(asset['id'])]
        put_rows(ws, add_row, 5, ('B', 'F', 'L'), [(r['year'], r['toYear'], r['amount']) for r in iv['add']])
        put_rows(ws, out_row, 5, ('B', 'F', 'L'), [(r['year'], r['toYear'], r['amount']) for r in iv['out']])


def recalc(xlsx, work):
    profile = Path(work) / 'lo-profile'
    out = Path(work) / 'out'
    subprocess.run(['soffice', f'-env:UserInstallation=file://{profile}', '--headless', '--convert-to', 'xlsx',
                    '--outdir', str(out), str(xlsx)], check=True, capture_output=True, timeout=300)
    return out / Path(xlsx).name


def read_cf(path):
    ws = openpyxl.load_workbook(path, data_only=True)['CF表']
    years, col = [], 4
    while isinstance(ws.cell(2, col).value, (int, float)) and ws.cell(3, col).value not in (None, ''):
        row = {'year': int(ws.cell(2, col).value)}
        for r, key in CF_ROWS.items():
            v = ws.cell(r, col).value
            row[key] = float(v) if isinstance(v, (int, float)) else 0.0
        years.append(row)
        col += 1
    return years


def main(src, work):
    work = Path(work)
    work.mkdir(parents=True, exist_ok=True)

    # 1) 置き換えた数式でも、サンプル入力でExcelの保存値を再現できるか
    wb = openpyxl.load_workbook(src)
    print(f'給与の数式を {patch_for_libreoffice(wb)} セル置き換え')
    probe = work / 'probe.xlsx'
    wb.save(probe)
    lo = read_cf(recalc(probe, work))
    excel = json.loads((ROOT / 'test/fixtures/excel-cf.json').read_text())['years']
    worst = max(abs(a[k] - b[k]) for a, b in zip(lo, excel) for k in b if k != 'year')
    if len(lo) != len(excel) or worst > 0.005:
        sys.exit(f'LibreOfficeの再計算がExcelの保存値と一致しません（最大差 {worst}）。照合に使えないので中止します')
    print(f'LibreOfficeの再計算はExcelの保存値と一致（{len(lo)}年、最大差 {worst:.2e}）')

    # 2) 各シナリオ
    for path in sorted((ROOT / 'test/scenarios').glob('*.json')):
        state = json.loads(path.read_text())
        wb = openpyxl.load_workbook(src)
        patch_for_libreoffice(wb)
        write_state(wb, state)
        x = work / f'{path.stem}.xlsx'
        wb.save(x)
        years = read_cf(recalc(x, work))
        out = ROOT / 'test/fixtures' / f'scenario-{path.stem}.json'
        out.write_text(json.dumps({'name': path.stem, 'state': state, 'years': years}, ensure_ascii=False, indent=1))
        print(f'{path.stem}: {len(years)}年分（{years[0]["year"]}〜{years[-1]["year"]}）')


if __name__ == '__main__':
    if len(sys.argv) not in (2, 3):
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2] if len(sys.argv) == 3 else tempfile.mkdtemp())
